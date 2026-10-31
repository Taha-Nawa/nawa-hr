"""
يولّد قيم «الإجابة الصحيحة» للمحرك من صيغ ملف الإدارة نفسه، بإعادة حسابها في LibreOffice، لا من الكود.

الفكرة: نأخذ ورقة «سجل الدوام اليومي» (بصيغها كما يستعملها الموظفون)، نكتب مكان أعمدة الحضور/الخروج/العودة/الانصراف
ونوع الإجازة مدخلات اختبار (لا بيانات أي موظف حقيقي)، ونضبط الإعدادات والعطل، ثم نترك LibreOffice يحسب الأعمدة K..T.
كل قيمة في ملفات test/fixtures هي ناتج صيغة إكسل لا ناتج المحرك، فلا يختبر المحرك نفسه.

الاستعمال (يحتاج LibreOffice، وملف الإدارة الحالي كقالب للصيغ فقط):
  python3 -I tools/recalc_golden.py --template "<ملف الإدارة 2026.xlsx>" --workdir /tmp/recalc \
          --old test/fixtures/march-2026-golden.json --out-old test/fixtures/march-2026-golden.json \
          --out-matrix test/fixtures/current-model-matrix.json

ملاحظة: قالب الصيغ يحتوي بيانات حقيقية، لذلك لا يُرفع إلى المستودع ولا تُنسخ منه إلا الصيغ.
"""
import argparse
import datetime as dt
import json
import os
import random
import shutil
import subprocess

import openpyxl

LOG = "سجل الدوام اليومي"
SET = "الإعدادات"
HOL = "العطل الرسمية"
FIRST_ROW = 4
DAYS_AR = ["الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت"]


def to_min(v):
    if v in (None, ""):
        return None
    h, m = str(v).split(":")[:2]
    return int(h) * 60 + int(m)


def hhmm(m):
    return None if m is None else f"{m // 60:02d}:{m % 60:02d}"


def as_time(m):
    return None if m is None else dt.time(m // 60, m % 60)


def recalc(template, workdir, tag, settings, holidays, rows):
    """rows: [{date, in, out, ret, dep, leave}] بالدقائق؛ يرجع [{...K..T...}] بحسب الترتيب."""
    os.makedirs(workdir, exist_ok=True)
    wb = openpyxl.load_workbook(template)
    s = wb[SET]
    s["C17"] = as_time(to_min(settings["start"]))
    s["C18"] = as_time(to_min(settings["end"]))
    s["C20"] = settings["toleranceMinutes"]
    s["C22"] = settings.get("monthlyExitCapHours", 6.0)
    for i, name in enumerate(DAYS_AR):
        s.cell(25 + i, 5).value = name
        s.cell(25 + i, 6).value = settings["weekdays"][i]
    for r in range(35, 43):  # أنواع الإجازات الثمانية في القالب
        pass
    leave_names = [s.cell(r, 5).value for r in range(35, 43)]
    for name in leave_names:
        s.cell(35 + leave_names.index(name), 6).value = "نعم" if settings["leaveTypes"].get(name) else "لا"
    h = wb[HOL]
    for r in range(4, 61):
        h.cell(r, 1).value = None
        h.cell(r, 2).value = None
    for i, d in enumerate(sorted(holidays)):
        h.cell(4 + i, 1).value = dt.datetime.strptime(d, "%Y-%m-%d")
        h.cell(4 + i, 2).value = "عطلة اختبار"

    log = wb[LOG]
    for i, row in enumerate(rows):
        r = FIRST_ROW + i
        log.cell(r, 5).value = as_time(row["in"])
        log.cell(r, 6).value = as_time(row["out"])
        log.cell(r, 7).value = as_time(row["ret"])
        log.cell(r, 8).value = as_time(row["dep"])
        log.cell(r, 9).value = row["leave"] or None
    src = os.path.join(workdir, f"{tag}.xlsx")
    wb.save(src)

    profile = os.path.join(workdir, "lo_profile")
    os.makedirs(os.path.join(profile, "user"), exist_ok=True)
    with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w", encoding="utf-8") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
            '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
            "<value>0</value></prop></item>\n</oor:items>\n"
        )
    out_dir = os.path.join(workdir, f"{tag}_out")
    shutil.rmtree(out_dir, ignore_errors=True)
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", out_dir, src],
        check=True, capture_output=True, timeout=300,
    )
    res = openpyxl.load_workbook(os.path.join(out_dir, f"{tag}.xlsx"), data_only=True)[LOG]
    num = lambda v: 0.0 if v in (None, "") else float(v)
    out = []
    for i, row in enumerate(rows):
        r = FIRST_ROW + i
        d = res.cell(r, 1).value
        assert isinstance(d, dt.datetime) and d.strftime("%Y-%m-%d") == row["date"], (r, d, row["date"])
        out.append({
            "dayType": res.cell(r, 4).value,
            "presenceHours": num(res.cell(r, 11).value),
            "exitHours": num(res.cell(r, 12).value),
            "netHours": num(res.cell(r, 13).value),
            "requiredHours": num(res.cell(r, 14).value),
            "approvedHours": round(num(res.cell(r, 15).value), 6),
            "lateMinutes": num(res.cell(r, 16).value),
            "earlyLeaveMinutes": num(res.cell(r, 17).value),
            "deficitHours": num(res.cell(r, 18).value),
            "overtimeHours": num(res.cell(r, 19).value),
            "status": res.cell(r, 20).value or "",
        })
    return out


# ---------------------------------------------------------------- 1) تحديث الـ fixture الحالي (نفس المدخلات، قيم إكسل الجديدة)
def refresh_old(template, workdir, old_path, out_path):
    fx = json.load(open(old_path, encoding="utf-8"))
    by = {}
    for m in fx["movements"]:
        if m["code"] == fx["employeeCode"]:
            by.setdefault(m["date"], {}).setdefault(m["type"], []).append(to_min(m["time"]))
    rows = []
    for d in fx["days"]:
        t = by.get(d["date"], {})
        man = d["manual"]
        ins, outs, rets, deps = t.get("حضور", []), t.get("خروج", []), t.get("عودة", []), t.get("انصراف", [])
        cin = to_min(man["in"]) if man["in"] else (min(ins) if ins else None)
        dep = to_min(man["dep"]) if man["dep"] else (max(deps) if deps else None)
        if man["out"] and man["ret"]:
            out_, ret = to_min(man["out"]), to_min(man["ret"])
        elif outs and len(outs) == len(rets):
            # القالب يحمل زوجاً واحداً (أول خروج/آخر عودة)؛ نكافئ مجموع الأزواج بزوج واحد له نفس المدة
            out_ = min(outs)
            ret = out_ + max(0, sum(rets) - sum(outs))
        else:
            out_, ret = (min(outs) if outs else None), (max(rets) if rets else None)
            if len(outs) != len(rets):
                ret = None  # خروج بلا عودة أو العكس: ساعات المغادرة صفر في الحالتين
        rows.append({"date": d["date"], "in": cin, "out": out_, "ret": ret, "dep": dep, "leave": d["leaveType"]})
    res = recalc(template, workdir, "old", fx["settings"], fx["holidays"], rows)
    changed = {}
    for d, r in zip(fx["days"], res):
        for k, v in r.items():
            if k == "dayType" and d["expected"][k] != v:
                raise SystemExit(f"نوع اليوم يختلف {d['date']}: {d['expected'][k]} ≠ {v}")
            if abs(float(d["expected"][k]) - float(v)) > 1e-9 if isinstance(v, float) else d["expected"][k] != v:
                changed.setdefault(k, 0)
                changed[k] += 1
            d["expected"][k] = v
    json.dump(fx, open(out_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("fixture الحالي: الحقول التي تغيّرت قيمتها وعدد الأيام:", changed or "لا شيء")


# ---------------------------------------------------------------- 2) مصفوفة اصطناعية بإعدادات الملف الحالي
def build_matrix(template, workdir, out_path):
    wv = openpyxl.load_workbook(template, data_only=True)
    s = wv[SET]
    weekdays = {s.cell(r, 5).value: s.cell(r, 6).value for r in range(25, 32)}
    settings = {
        "start": hhmm(s["C17"].value.hour * 60 + s["C17"].value.minute),
        "end": hhmm(s["C18"].value.hour * 60 + s["C18"].value.minute),
        "toleranceMinutes": int(s["C20"].value),
        "monthlyExitCapHours": float(s["C22"].value),
        "weekdays": [weekdays[d] for d in DAYS_AR],
        "leaveTypes": {s.cell(r, 5).value: s.cell(r, 6).value == "نعم" for r in range(35, 43)},
    }
    holidays = sorted(
        r[0].strftime("%Y-%m-%d") for r in wv[HOL].iter_rows(min_row=4, values_only=True) if isinstance(r[0], dt.datetime)
    )
    rnd = random.Random(20261031)
    start, end, tol = to_min(settings["start"]), to_min(settings["end"]), settings["toleranceMinutes"]
    leaves = list(settings["leaveTypes"])
    rows = []
    day = dt.date(2026, 1, 1)
    while day.year == 2026:
        kind = rnd.choices(
            ["normal", "late_in_tol", "late_out_tol", "early_bird", "exit_return", "exit_no_return", "no_in", "no_dep",
             "absent", "leave_only", "leave_attend", "overtime", "early_leave", "boundary"],
            weights=[18, 10, 8, 6, 8, 3, 2, 3, 8, 5, 4, 4, 5, 6],
        )[0]
        cin = out_ = ret = dep = None
        leave = ""
        if kind == "normal":
            cin, dep = start + rnd.randint(-3, 0), end + rnd.randint(0, 4)
        elif kind == "late_in_tol":
            cin, dep = start + rnd.randint(1, tol), end + rnd.randint(-2, 6)
        elif kind == "late_out_tol":
            cin, dep = start + rnd.randint(tol + 1, 70), end + rnd.randint(-5, 10)
        elif kind == "early_bird":
            cin, dep = start - rnd.randint(5, 50), end + rnd.randint(0, 20)
        elif kind == "exit_return":
            cin, dep = start + rnd.randint(-5, 25), end + rnd.randint(-10, 15)
            out_ = rnd.randint(cin + 30, dep - 90)
            ret = out_ + rnd.randint(5, 80)
        elif kind == "exit_no_return":
            cin, dep = start + rnd.randint(0, 12), end + rnd.randint(-10, 5)
            out_ = rnd.randint(cin + 30, dep - 60)
        elif kind == "no_in":
            dep = end + rnd.randint(-30, 10)
        elif kind == "no_dep":
            cin = start + rnd.randint(-5, 30)
        elif kind == "leave_only":
            leave = rnd.choice(leaves)
        elif kind == "leave_attend":
            leave = rnd.choice(leaves)
            cin, dep = start + rnd.randint(-10, 40), end + rnd.randint(-120, 90)
        elif kind == "overtime":
            cin, dep = start + rnd.randint(-20, 5), end + rnd.randint(30, 240)
        elif kind == "early_leave":
            cin, dep = start + rnd.randint(-5, 20), end - rnd.randint(10, 150)
        elif kind == "boundary":  # حدود السماحية بالضبط
            cin, dep = start + rnd.choice([tol - 1, tol, tol + 1, 0, 1]), end + rnd.choice([0, 1, -1])
        # أيام العطل والأسبوع تحصل على نفس الأنماط عشوائياً فنختبر «عطلة فيها دوام» و«إجازة في عطلة»
        rows.append({"date": day.isoformat(), "in": cin, "out": out_, "ret": ret, "dep": dep, "leave": leave})
        day += dt.timedelta(days=1)
    res = recalc(template, workdir, "matrix", settings, holidays, rows)
    days = []
    for row, r in zip(rows, res):
        mv = []
        for typ, key in (("حضور", "in"), ("خروج", "out"), ("عودة", "ret"), ("انصراف", "dep")):
            if row[key] is not None:
                mv.append({"type": typ, "time": hhmm(row[key])})
        days.append({"date": row["date"], "movements": mv, "leaveType": row["leave"], "expected": r})
    json.dump(
        {"source": "صيغ ملف الإدارة 2026 (أعيد حسابها في LibreOffice على مدخلات اصطناعية)", "settings": settings,
         "holidays": holidays, "days": days},
        open(out_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1,
    )
    import collections
    print("المصفوفة:", len(days), "يوماً؛ الحالات:", dict(collections.Counter(d["expected"]["status"] for d in days)))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--workdir", required=True)
    ap.add_argument("--old")
    ap.add_argument("--out-old")
    ap.add_argument("--out-matrix")
    a = ap.parse_args()
    if a.old:
        refresh_old(a.template, a.workdir, a.old, a.out_old or a.old)
    if a.out_matrix:
        build_matrix(a.template, a.workdir, a.out_matrix)
