"""
Extracts a "golden" fixture from the management Excel workbook so the JS engine can be
verified against the exact numbers Excel produced (cached values).

Usage:  python3 -I tools/extract_golden.py <workbook.xlsx> <out.json>

Reads: الإعدادات (settings), العطل الرسمية (holidays), الحركات الواردة (movements),
       سجل الدوام اليومي (daily log: manual overrides, leave type, and Excel's computed outputs).
"""
import datetime as dt
import json
import sys

import openpyxl


def to_minutes(v):
    """Normalise an Excel time-ish cell to whole minutes since midnight (or None)."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)):
        return round((v % 1) * 1440)
    s = str(v).strip()
    if not s:
        return None
    h, m = s.split(":")[:2]
    return int(h) * 60 + int(m)


def hhmm(minutes):
    return None if minutes is None else f"{minutes // 60:02d}:{minutes % 60:02d}"


def num(v):
    return 0 if v in (None, "") else float(v)


def main(src, out):
    wf = openpyxl.load_workbook(src, data_only=False)
    wv = openpyxl.load_workbook(src, data_only=True)

    s = wv["الإعدادات"]
    weekdays = {s.cell(r, 5).value: s.cell(r, 6).value for r in range(25, 32)}
    leave_types = {s.cell(r, 5).value: s.cell(r, 6).value == "نعم" for r in range(35, 43)}
    settings = {
        "start": hhmm(to_minutes(s["C17"].value)),
        "end": hhmm(to_minutes(s["C18"].value)),
        "toleranceMinutes": int(s["C20"].value),
        "monthlyExitCapHours": float(s["C22"].value),
        # Excel WEEKDAY(date,1): 1=Sunday ... 7=Saturday
        "weekdays": [weekdays[d] for d in
                     ["الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت"]],
        "leaveTypes": leave_types,
    }
    employee_code = s["C6"].value

    h = wv["العطل الرسمية"]
    holidays = [h.cell(r, 1).value.strftime("%Y-%m-%d")
                for r in range(4, h.max_row + 1) if isinstance(h.cell(r, 1).value, dt.datetime)]

    m = wv["الحركات الواردة"]
    movements = []
    for r in range(4, m.max_row + 1):
        if m.cell(r, 1).value in (None, ""):
            continue
        movements.append({
            "id": m.cell(r, 1).value,
            "code": m.cell(r, 3).value,
            "type": m.cell(r, 5).value,
            "date": str(m.cell(r, 6).value)[:10],
            "time": hhmm(to_minutes(m.cell(r, 7).value)),
        })

    d = wv["سجل الدوام اليومي"]
    days = []
    for r in range(4, d.max_row + 1):
        date = d.cell(r, 1).value
        if not isinstance(date, dt.datetime):
            continue
        days.append({
            "date": date.strftime("%Y-%m-%d"),
            "manual": {
                "in": hhmm(to_minutes(d.cell(r, 5).value)),
                "out": hhmm(to_minutes(d.cell(r, 6).value)),
                "ret": hhmm(to_minutes(d.cell(r, 7).value)),
                "dep": hhmm(to_minutes(d.cell(r, 8).value)),
            },
            "leaveType": d.cell(r, 9).value or "",
            "expected": {
                "dayType": d.cell(r, 4).value,
                "checkIn": hhmm(to_minutes(d.cell(r, 11).value)),
                "exitFirst": hhmm(to_minutes(d.cell(r, 12).value)),
                "returnLast": hhmm(to_minutes(d.cell(r, 13).value)),
                "departure": hhmm(to_minutes(d.cell(r, 14).value)),
                "source": d.cell(r, 15).value,
                "presenceHours": num(d.cell(r, 16).value),
                "exitHours": num(d.cell(r, 17).value),
                "netHours": num(d.cell(r, 18).value),
                "requiredHours": num(d.cell(r, 19).value),
                "approvedHours": num(d.cell(r, 20).value),
                "lateMinutes": num(d.cell(r, 21).value),
                "earlyLeaveMinutes": num(d.cell(r, 22).value),
                "deficitHours": num(d.cell(r, 23).value),
                "overtimeHours": num(d.cell(r, 24).value),
                "status": d.cell(r, 25).value or "",
                "alert": d.cell(r, 26).value or "",
            },
        })

    json.dump({"source": "نموذج تشغيلي مختبَر - مارس 2026", "employeeCode": employee_code,
               "settings": settings, "holidays": holidays, "movements": movements, "days": days},
              open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"days={len(days)} movements={len(movements)} holidays={len(holidays)} -> {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
